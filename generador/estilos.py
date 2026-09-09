"""Estilos y utilidades de documento de la guía DOCX."""
import os

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .config import (AZUL, AZUL_CLARO, ELABORO, FUENTE_CODIGO, FUENTE_NORMAL,
                     GRIS_CLARO)


def leer(ej_dir, ruta_relativa):
    ruta = os.path.join(ej_dir, ruta_relativa)
    with open(ruta, "r", encoding="utf-8") as f:
        return f.read()


def configurar_estilos(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FUENTE_NORMAL
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "es-CO")
    rpr.append(lang)
    tamanos = {1: 16, 2: 13, 3: 12}
    for nivel in range(1, 4):
        nombre = f"Heading {nivel}"
        if nombre in doc.styles:
            st = doc.styles[nombre]
            st.font.color.rgb = AZUL
            st.font.size = Pt(tamanos.get(nivel, 12))
    doc.core_properties.title = "Guía de Aprendizaje — Testing & QA (SENA ADSO)"
    doc.core_properties.subject = (
        "Tecnólogo en Análisis y Desarrollo de Software — adaptación didáctica"
    )
    doc.core_properties.author = ELABORO
    doc.core_properties.comments = (
        "Adaptación didáctica local; validar datos y formato de entrega con el instructor."
    )


def sombrear(celda, color=AZUL_CLARO):
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def tabla(documento, cabeceras, filas, ancho_cm=None):
    t = documento.add_table(rows=1 + len(filas), cols=len(cabeceras))
    t.style = "Table Grid"
    for idx, h in enumerate(cabeceras):
        c = t.rows[0].cells[idx]
        c.text = h
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(10)
        sombrear(c)
    for fila_idx, fila in enumerate(filas, start=1):
        for col_idx, valor in enumerate(fila):
            celda = t.rows[fila_idx].cells[col_idx]
            celda.text = valor
            for p in celda.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    if ancho_cm:
        for row in t.rows:
            for idx, ancho in enumerate(ancho_cm):
                row.cells[idx].width = Cm(ancho)
    return t


def agregar_codigo(doc, codigo):
    for linea in codigo.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        pPr = p._element.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), GRIS_CLARO)
        pPr.append(shd)
        run = p.add_run(linea if linea else " ")
        run.font.name = FUENTE_CODIGO
        run.font.size = Pt(9)


def agregar_pasos(doc, titulo, pasos):
    doc.add_heading(titulo, level=3)
    for paso in pasos:
        par = doc.add_paragraph()
        run_num = par.add_run(f"Paso {paso['num']}: {paso['titulo']}. ")
        run_num.font.bold = True
        par.add_run(paso["desc"])
        if paso.get("cmd"):
            p_cmd = doc.add_paragraph()
            p_cmd.paragraph_format.left_indent = Cm(0.8)
            p_cmd.paragraph_format.space_before = Pt(1)
            p_cmd.paragraph_format.space_after = Pt(2)
            run_cmd = p_cmd.add_run(f"> {paso['cmd']}")
            run_cmd.font.name = FUENTE_CODIGO
            run_cmd.font.size = Pt(9)
            run_cmd.font.color.rgb = AZUL
        if paso.get("tip"):
            p_tip = doc.add_paragraph()
            p_tip.paragraph_format.left_indent = Cm(0.8)
            p_tip.paragraph_format.space_before = Pt(0)
            p_tip.paragraph_format.space_after = Pt(2)
            run_tip = p_tip.add_run(f"Regla de oro: {paso['tip']}")
            run_tip.font.size = Pt(9)
            run_tip.font.italic = True


def agregar_componentes(doc, filas):
    tabla(doc, ["Componente", "Orientación"], filas, ancho_cm=[5.2, 10.0])


def agregar_campo_pagina(doc):
    seccion = doc.sections[0]
    pie = seccion.footer
    p = pie.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Guía de Aprendizaje — Testing & QA · SENA ADSO · Adaptación didáctica · Página "
    )
    r.font.size = Pt(8)
    run = p.add_run()
    run.font.size = Pt(8)
    _agregar_campo(run, " PAGE ")
    r2 = p.add_run(" de ")
    r2.font.size = Pt(8)
    run2 = p.add_run()
    run2.font.size = Pt(8)
    _agregar_campo(run2, " NUMPAGES ")


def _agregar_campo(run, instruccion):
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruccion
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
