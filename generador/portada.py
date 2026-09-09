"""Portada institucional de la guía."""
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .config import (AZUL, CENTRO, FASE_PROYECTO, GRIS, PROGRAMA, REGIONAL,
                     VERSION, fecha_es_co)


def portada(doc):
    for _ in range(3):
        doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PROCESO DE GESTIÓN DE FORMACIÓN PROFESIONAL INTEGRAL")
    r.font.size = Pt(11)
    r.font.color.rgb = GRIS
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FORMATO GUÍA DE APRENDIZAJE")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = AZUL
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SERVICIO NACIONAL DE APRENDIZAJE")
    r.font.size = Pt(18)
    r.font.bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{REGIONAL} · {CENTRO}")
    r.font.size = Pt(11)
    r.font.color.rgb = GRIS
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("GUÍA 05")
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = AZUL
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Verificación y aseguramiento de la calidad del software mediante "
        "pruebas automatizadas (Testing & QA)"
    )
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = AZUL
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Programa {PROGRAMA} - ADSO")
    r.font.size = Pt(14)
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        f"Fase del proyecto: {FASE_PROYECTO} — Calidad de Software\n"
        f"Adaptación didáctica · Versión {VERSION} · {fecha_es_co()}"
    )
    r.font.size = Pt(10)
    r.font.color.rgb = GRIS
    doc.add_page_break()
