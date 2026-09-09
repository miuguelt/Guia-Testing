"""Sección 1: identificación institucional de la guía de aprendizaje."""
from docx.shared import Cm

from .config import (ACTIVIDAD_PROYECTO, CODIGO_PROGRAMA, COMPETENCIA,
                     DURACION_TOTAL, FASE_PROYECTO, MODALIDAD, PROGRAMA,
                     PROYECTO_FORMATIVO, RESULTADOS_APRENDIZAJE, VERSION)


def identificacion(doc):
    doc.add_heading("1. IDENTIFICACIÓN DE LA GUÍA DE APRENDIZAJE", level=1)
    campo(doc, "Denominación del Programa de Formación", PROGRAMA)
    campo(doc, "Código del Programa de Formación", CODIGO_PROGRAMA)
    campo(doc, "Nombre del Proyecto Formativo", PROYECTO_FORMATIVO)
    campo(doc, "Fase del Proyecto", FASE_PROYECTO)
    campo(doc, "Actividad de Proyecto Formativo", ACTIVIDAD_PROYECTO)
    campo(doc, "Competencia", COMPETENCIA)
    p = doc.add_paragraph()
    r = p.add_run("Resultados de Aprendizaje: ")
    r.font.bold = True
    for i, ra in enumerate(RESULTADOS_APRENDIZAJE, 1):
        part = doc.add_paragraph(f"{i}. {ra}")
        part.paragraph_format.left_indent = Cm(0.75)
    campo(doc, "Duración de la Guía de Aprendizaje", DURACION_TOTAL)
    campo(doc, "Modalidad", MODALIDAD)
    campo(doc, "Versión institucional", f"GFPI-F-135 — {VERSION}")
    doc.add_page_break()


def campo(doc, label, valor):
    p = doc.add_paragraph()
    r = p.add_run(f"{label}: ")
    r.font.bold = True
    p.add_run(valor)
