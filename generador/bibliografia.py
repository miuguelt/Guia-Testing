"""Sección 8: referentes bibliográficos."""
from docx.shared import Pt


def bibliografia(doc):
    doc.add_heading("8. REFERENTES BIBLIOGRÁFICOS", level=1)
    doc.add_paragraph(
        "Referentes curriculares, técnicos y de apoyo consultados para la "
        "elaboración de la guía:"
    )
    referencias = [
        "Servicio Nacional de Aprendizaje - SENA. (2021). Programa de "
        "formación titulada: Análisis y Desarrollo de Software, código "
        "228118.",
        "Servicio Nacional de Aprendizaje - SENA. (s. f.). Formato Guía de "
        "Aprendizaje GFPI-F-135.",
        "Institute of Electrical and Electronics Engineers. (2008). IEEE "
        "829: Standard for Software and System Test Documentation.",
        "International Organization for Standardization. (2011). ISO/IEC "
        "25010: Systems and software Quality Requirements and Evaluation "
        "(SQuaRE).",
        "Beck, K. (2003). Test-Driven Development: By Example. "
        "Addison-Wesley.",
        "Cohn, M. (2009). Succeeding with Agile. Addison-Wesley.",
        "North, D. (2006). Introducing BDD. Better Software Magazine.",
        "Playwright. (2024). Documentación oficial. https://playwright.dev/docs",
        "PyTest. (2024). Documentación oficial. https://docs.pytest.org/",
        "Jest. (2024). Documentación oficial. https://jestjs.io/docs",
        "JUnit 5. (2024). Documentación oficial. https://junit.org/junit5/docs/",
        "GitHub Actions. (2024). Documentación oficial. "
        "https://docs.github.com/actions",
        "OWASP. (2024). Web Security Testing Guide. "
        "https://owasp.org/www-project-web-security-testing-guide/",
        "Microsoft. (2024). Playwright Best Practices.",
    ]
    for referencia in referencias:
        p = doc.add_paragraph(referencia, style="List Bullet")
        for run in p.runs:
            run.font.size = Pt(11)
    doc.add_page_break()
